#!/usr/bin/env -S uv run --script
# /// script
# requires-python = ">=3.10"
# dependencies = ["openpyxl"]
# ///
"""Fetch a Google Doc/Sheet/Slides/Drive file by link and make it agent-readable.

Order: rclone remote `gdrive:` (override with GDOC_RCLONE_REMOTE) (read-only OAuth, sees files shared with the user),
then the public export URL (works when the file is "anyone with the link").
Sheets are also split into one CSV per tab.
"""
import argparse, csv, os, re, subprocess, sys, urllib.request
from pathlib import Path

REMOTE = os.environ.get("GDOC_RCLONE_REMOTE", "gdrive") + ":"
KINDS = {"spreadsheets": "sheet", "document": "doc", "presentation": "slides"}
PUBLIC_FMT = {"sheet": "xlsx", "doc": "md", "slides": "pptx"}


def parse(link: str):
    m = re.search(r"docs\.google\.com/(spreadsheets|document|presentation)/d/([\w-]+)", link)
    if m:
        return KINDS[m.group(1)], m.group(2)
    m = re.search(r"(?:/file/d/|[?&]id=)([\w-]+)", link)
    if m:
        return "file", m.group(1)
    if re.fullmatch(r"[\w-]{20,}", link):
        return "file", link
    sys.exit(f"无法识别的链接: {link}")


def via_rclone(fid: str, out: Path) -> Path | None:
    before = set(out.iterdir())
    r = subprocess.run(
        ["rclone", "backend", "copyid", REMOTE, fid, f"{out}/",
         "--drive-export-formats", "xlsx,md,pptx,pdf", "-q"],
        capture_output=True, text=True)
    new = set(out.iterdir()) - before
    if r.returncode == 0 and new:
        return new.pop()
    print(f"[rclone 失败] {r.stderr.strip()[:300]}", file=sys.stderr)
    return None


def via_public(kind: str, fid: str, out: Path) -> Path | None:
    if kind == "file":
        url = f"https://drive.google.com/uc?export=download&id={fid}"
        dest = out / fid
    else:
        fmt = PUBLIC_FMT[kind]
        path = {"sheet": "spreadsheets", "doc": "document", "slides": "presentation"}[kind]
        url = f"https://docs.google.com/{path}/d/{fid}/export?format={fmt}"
        dest = out / f"{fid}.{fmt}"
    try:
        with urllib.request.urlopen(url, timeout=60) as resp:
            if "text/html" in resp.headers.get("Content-Type", ""):
                raise RuntimeError("返回了登录页，文件不是公开的")
            dest.write_bytes(resp.read())
        return dest
    except Exception as e:
        print(f"[公开链接失败] {e}", file=sys.stderr)
        return None


def cell(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return int(v)
    return v


def xlsx_to_csv(path: Path) -> list[Path]:
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    outs = []
    for ws in wb.worksheets:
        name = re.sub(r'[\\/:*?"<>|]', "_", ws.title)
        dest = path.with_name(f"{path.stem}__{name}.csv")
        with dest.open("w", newline="", encoding="utf-8") as f:
            w = csv.writer(f)
            for row in ws.iter_rows(values_only=True):
                w.writerow([cell(v) for v in row])
        outs.append(dest)
    return outs


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("link", help="Google 文档/表格/幻灯片/Drive 链接或文件 ID")
    ap.add_argument("--out", default="/tmp/gdoc", help="输出目录 (默认 /tmp/gdoc)")
    a = ap.parse_args()
    kind, fid = parse(a.link)
    out = Path(a.out) / fid
    out.mkdir(parents=True, exist_ok=True)
    got = via_rclone(fid, out) or via_public(kind, fid, out)
    if not got:
        sys.exit("下载失败：确认链接已分享给已授权的谷歌账号，或让对方设为「知道链接的人可查看」。")
    print(got)
    if got.suffix == ".xlsx":
        for p in xlsx_to_csv(got):
            print(p)


if __name__ == "__main__":
    main()
